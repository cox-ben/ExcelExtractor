import io
import openpyxl
from excel_extractor.engine import QuoteExtractorEngine


def test_mpc_header_at_row_31_with_index_row_and_top_metadata():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MPC"

    # Metadata at top of MPC sheet (rows 1-10) with labels that could mimic headers
    ws["A1"] = "Molded Precision Components"
    ws["A2"] = "Quotation Summary"
    ws["B4"] = "Customer: Test Corp"
    ws["B5"] = "Quote Date: 2026-09-23"
    ws["B6"] = "Quote Number:"
    ws["C6"] = "Q4509"
    ws["D6"] = "Part Number:"
    ws["E6"] = "GENERAL-MOLDING"

    # Row 31: The REAL table header
    ws["A31"] = "Line Item"
    ws["B31"] = "MPC RFQ #"
    ws["C31"] = "Part No."
    ws["D31"] = "Description"
    ws["E31"] = "Annual Volume"

    # Row 32: Blank line

    # Row 33: Column index row (1, 2, 3, 4 across columns B-E)
    ws["B33"] = 1
    ws["C33"] = 2
    ws["D33"] = 3
    ws["E33"] = 4

    # Row 34: Data row 1
    ws["A34"] = 1
    ws["B34"] = "Q4509-1"
    ws["C34"] = "29729-BRG002"
    ws["D34"] = "BEARING HOUSING"
    ws["E34"] = 2000000

    # Row 35: Data row 2
    ws["A35"] = 2
    ws["B35"] = "Q4509-2"
    ws["C35"] = "40193-PPT002"
    ws["D35"] = "POPPET"
    ws["E35"] = 400000

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    engine = QuoteExtractorEngine()
    doc = engine.extract_file(buf, filename="Q4509_test.xlsx")

    assert doc.total_items == 2, f"Expected 2 items, got {doc.total_items}"
    assert doc.items[0].quote_item_number == "Q4509-1"
    assert doc.items[0].part_number == "29729-BRG002"
    assert doc.items[0].description == "BEARING HOUSING"
    assert doc.items[0].annual_volume == 2000000

    assert doc.items[1].quote_item_number == "Q4509-2"
    assert doc.items[1].part_number == "40193-PPT002"
    assert doc.items[1].description == "POPPET"
    assert doc.items[1].annual_volume == 400000


def test_mpc_header_at_row_40_with_spacer_and_float_index_row():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MPC"

    # Metadata scattered in top 30 rows
    ws["A2"] = "Customer: General Motors"
    ws["B5"] = "Quote Date: 2026-09-23"
    ws["B12"] = "Quote Number: Q5120"
    ws["D12"] = "Part Number: SAMPLE"

    # Row 40: Table Header
    ws["A40"] = "Item #"
    ws["B40"] = "Quote #"
    ws["C40"] = "Customer Part #"
    ws["D40"] = "Part Description"
    ws["E40"] = "Cycle Time (sec)"
    ws["F40"] = "Annual Volume"

    # Row 41: Blank spacer line
    # Row 42: Float column index row
    ws["B42"] = 1.0
    ws["C42"] = 2.0
    ws["D42"] = 3.0
    ws["E42"] = 4.0
    ws["F42"] = 5.0

    # Row 43: Another blank spacer line
    # Row 44: Item 1
    ws["A44"] = 1
    ws["B44"] = "Q5120-1"
    ws["C44"] = "98234-A1"
    ws["D44"] = "VALVE SEAT"
    ws["E44"] = 18.5
    ws["F44"] = 500000

    # Row 45: Item 2
    ws["A45"] = 2
    ws["B45"] = "Q5120-2"
    ws["C45"] = "98234-A2"
    ws["D45"] = "VALVE PLUG"
    ws["E45"] = 22.0
    ws["F45"] = 750000

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    engine = QuoteExtractorEngine()
    doc = engine.extract_file(buf, filename="Q5120_test.xlsx")

    assert doc.total_items == 2, f"Expected 2 items, got {doc.total_items}"
    assert doc.items[0].quote_item_number == "Q5120-1"
    assert doc.items[0].part_number == "98234-A1"
    assert doc.items[0].cycle_time_sec == 18.5
    assert doc.items[0].annual_volume == 500000

    assert doc.items[1].quote_item_number == "Q5120-2"
    assert doc.items[1].part_number == "98234-A2"
    assert doc.items[1].cycle_time_sec == 22.0
    assert doc.items[1].annual_volume == 750000


def test_checklist_with_questions_and_letters_ignored():
    """Verify that RFQ checklist/evaluation tables with questions & letters are never treated as quote items."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MPC"

    # Qualification checklist at row 10
    ws["A10"] = "Item #"
    ws["B10"] = "Part Review"
    ws["A11"] = 1; ws["B11"] = "Strategic Fit for MPC?"
    ws["A12"] = 2; ws["B12"] = "Customer Awarded?"
    ws["A13"] = 3; ws["B13"] = "RFQ Package Complete?"
    ws["A14"] = 4; ws["B14"] = "MPC Current Material"
    ws["A15"] = 5; ws["B15"] = "Priority Grade"
    ws["A16"] = 6; ws["B16"] = "T"
    ws["A17"] = 7; ws["B17"] = "A"
    ws["A18"] = 8; ws["B18"] = "B"
    ws["A19"] = 9; ws["B19"] = "C"
    ws["A20"] = 10; ws["B20"] = "D"

    # Real quote table starting at row 31
    ws["A31"] = "Line Item"
    ws["B31"] = "MPC RFQ #"
    ws["C31"] = "Part No."
    ws["D31"] = "Description"
    ws["E31"] = "Annual Volume"

    ws["B33"] = 2; ws["C33"] = 3; ws["D33"] = 4; ws["E33"] = 5

    ws["A34"] = 1; ws["B34"] = "Q4509-1"; ws["C34"] = "29729-BRG002"; ws["D34"] = "BEARING HOUSING"; ws["E34"] = 2000000
    ws["A35"] = 2; ws["B35"] = "Q4509-2"; ws["C35"] = "40193-PPT002"; ws["D35"] = "POPPET"; ws["E35"] = 400000
    ws["A36"] = 3; ws["B36"] = "Q4509-3"; ws["C36"] = "40193-PPT003"; ws["D36"] = "POPPET - CV"; ws["E36"] = 400000
    ws["A37"] = 4; ws["B37"] = "Q4509-4"; ws["C37"] = "40195-PPT002"; ws["D37"] = "POPPET"; ws["E37"] = 4000000

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    engine = QuoteExtractorEngine()
    doc = engine.extract_file(buf, filename="Q4509_real.xlsx")

    assert doc.total_items == 4, f"Expected 4 items, got {doc.total_items}"
    assert [it.quote_item_number for it in doc.items] == ["Q4509-1", "Q4509-2", "Q4509-3", "Q4509-4"]
    assert [it.part_number for it in doc.items] == [
        "29729-BRG002",
        "40193-PPT002",
        "40193-PPT003",
        "40195-PPT002",
    ]
    assert [it.description for it in doc.items] == [
        "BEARING HOUSING",
        "POPPET",
        "POPPET - CV",
        "POPPET",
    ]
    assert [it.annual_volume for it in doc.items] == [2000000, 400000, 400000, 4000000]
