"""
Unit tests for extraction strategies, anchor locator, and engine coordinator.
"""

from __future__ import annotations

from pathlib import Path
import pytest

from excel_extractor.config import ConfigManager, ExtractorConfig
from excel_extractor.engine import QuoteExtractorEngine
from excel_extractor.locator import (
    AnchorLocator,
    clean_cycle_time_value,
    clean_part_number_string,
    is_drawing_anchor,
    is_drawing_value,
    normalize_text,
)
from excel_extractor.models import QuoteDocument, QuoteItem
from excel_extractor.reader import Workbook, Worksheet, load_workbook
from excel_extractor.strategies.grid_matrix import GridMatrixStrategy
from excel_extractor.strategies.multi_tab import MultiTabStrategy
from excel_extractor.strategies.single_item import SingleItemStrategy
from tests.fixtures import generator


class TestLocatorUtilities:
    """Tests for text normalization, cleaning, and disambiguation routines."""

    def test_normalize_text(self):
        assert normalize_text("  Part Number:  ") == "part number"
        assert normalize_text("  Quote # : ; - ") == "quote #"
        assert normalize_text("Customer P/N:\u00a0") == "customer p/n"
        assert normalize_text(None) == ""

    def test_is_drawing_anchor(self):
        assert is_drawing_anchor("Drawing Number:")
        assert is_drawing_anchor("DWG #:")
        assert is_drawing_anchor("dwg no")
        assert not is_drawing_anchor("Part Number:")
        assert not is_drawing_anchor("Customer Part:")

    def test_is_drawing_value(self):
        assert is_drawing_value("DWG-8841-B")
        assert is_drawing_value("DWG # 4022-A")
        assert is_drawing_value("DRAWING-100")
        assert not is_drawing_value("2202046213000")
        assert not is_drawing_value("MPC-7011-A")

    def test_clean_part_number_string(self):
        # Numeric integer conversion
        assert clean_part_number_string(2202046213000) == "2202046213000"
        # Numeric float with integer value
        assert clean_part_number_string(2202046213000.0) == "2202046213000"
        # Stripping trailing drawing number annotations
        assert clean_part_number_string("3042605C (DWG 4022-A)") == "3042605C"
        assert clean_part_number_string("3042605C [drawing 999]") == "3042605C"
        assert clean_part_number_string("858007801000\n(P0900612)") == "858007801000"
        assert clean_part_number_string("858007802000 (P0900712)") == "858007802000"
        assert clean_part_number_string("2202046201000 [P09000410]") == "2202046201000"
        # Non-breaking space
        assert clean_part_number_string("  3042605C-REV1; \u00a0") == "3042605C-REV1"

    def test_clean_cycle_time_value(self):
        assert clean_cycle_time_value(24.5) == 24.5
        assert clean_cycle_time_value(" 26.5 sec ") == 26.5
        assert clean_cycle_time_value("18.2s") == 18.2
        assert clean_cycle_time_value("35 seconds") == 35.0
        assert clean_cycle_time_value("#N/A") is None
        assert clean_cycle_time_value("") is None
        assert clean_cycle_time_value(None) is None


class TestAnchorLocatorSpatialResolution:
    """Tests for AnchorLocator merged-cell boundary traversal and spacer skipping."""

    def test_resolve_target_cell_merged_boundary(self):
        ws = Worksheet("MergedTest")
        # Label merged B6:C6
        ws.merge_cells("B6:C6")
        ws["B6"] = "Customer Part Number:"
        # Value merged D6:F6
        ws.merge_cells("D6:F6")
        ws["D6"] = "3042605C"

        locator = AnchorLocator()
        # Offset +1 col to the right of merged label should land on D6 (since C6 is max_col of label)
        cell, coord = locator.resolve_target_cell(ws, anchor_row=6, anchor_col=2, offset_row=0, offset_col=1)
        assert coord == "D6"
        assert cell.value == "3042605C"

    def test_resolve_target_cell_spacer_advancement(self):
        ws = Worksheet("SpacerTest")
        # Label merged B14:C14
        ws.merge_cells("B14:C14")
        ws["B14"] = "Cycle Time (Seconds):"
        ws["D14"] = ""  # empty spacer
        ws["E14"] = 28.0

        locator = AnchorLocator()
        cell, coord = locator.resolve_target_cell(ws, anchor_row=14, anchor_col=2, offset_row=0, offset_col=1)
        assert coord == "E14"
        assert cell.value == 28.0


class TestExtractionStrategies:
    """Tests for SingleItemStrategy, MultiTabStrategy, and GridMatrixStrategy."""

    def test_single_item_strategy_modern(self, single_modern_quote_file: Path):
        wb = load_workbook(single_modern_quote_file)
        strat = SingleItemStrategy()
        assert strat.can_handle(wb)
        items = strat.extract(wb)
        assert len(items) == 1
        assert items[0].quote_item_number == "Q1001-1"
        assert items[0].part_number == "2202046213000"
        assert items[0].cycle_time_sec == 24.5

    def test_single_item_strategy_inferred_suffix(self, legacy_raw_quote_file: Path):
        wb = load_workbook(legacy_raw_quote_file)
        strat = SingleItemStrategy()
        assert strat.can_handle(wb)
        items = strat.extract(wb)
        assert len(items) == 1
        assert items[0].quote_item_number == "Q3050-1"
        assert any("inferred" in w.lower() for w in items[0].warnings)

    def test_multi_tab_strategy_filtering_and_extraction(self, multitab_quote_file: Path):
        wb = load_workbook(multitab_quote_file)
        strat = MultiTabStrategy()
        assert strat.can_handle(wb)
        items = strat.extract(wb)
        assert len(items) == 3
        nums = [it.quote_item_number for it in items]
        assert nums == ["Q2005-1", "Q2005-2", "Q2005-3"]
        assert "Summary" not in nums
        assert "Lookups" not in nums

    def test_grid_matrix_strategy_header_and_data_rows(self, grid_matrix_quote_file: Path):
        wb = load_workbook(grid_matrix_quote_file)
        strat = GridMatrixStrategy()
        assert strat.can_handle(wb)
        items = strat.extract(wb)
        assert len(items) == 3
        assert [it.quote_item_number for it in items] == ["Q4000-1", "Q4000-2", "Q4000-3"]
        assert [it.part_number for it in items] == ["MED-10492", "MED-10493", "MED-10494"]
        assert [it.cycle_time_sec for it in items] == [14.5, 22.0, 9.8]

    def test_grid_matrix_part_id_vs_part_number_disambiguation(self):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "QInfo"
        ws.cell(row=1, column=2, value=1)
        ws.cell(row=2, column=1, value="PART ID")
        ws.cell(row=2, column=2, value="Item#")
        ws.cell(row=2, column=3, value="Quote Number")
        ws.cell(row=2, column=4, value="Drawing #")
        ws.cell(row=2, column=5, value="Part Number")
        ws.cell(row=2, column=6, value="Desc")
        ws.cell(row=2, column=7, value="Volume (EAU)")
        ws.cell(row=3, column=1, value="3142")
        ws.cell(row=3, column=2, value=10)
        ws.cell(row=3, column=3, value="Q4400-10")
        ws.cell(row=3, column=4, value="P09000810")
        ws.cell(row=3, column=5, value="858008003000\n(P09000810)")
        ws.cell(row=3, column=6, value="Spring Control 3 AU5\n(S233P / S233G)")
        ws.cell(row=3, column=7, value=69745)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        engine = QuoteExtractorEngine()
        doc = engine.extract_file(buf, filename="Q4400_test.xlsx")
        assert doc.total_items == 1
        item = doc.items[0]
        assert item.quote_item_number == "Q4400-10"
        assert item.part_number == "858008003000"
        assert item.drawing_number == "P09000810"
        assert item.annual_volume == 69745
        assert "Spring Control" in (item.description or "")


class TestEngineCoordinator:
    """Tests for QuoteExtractorEngine initialization and operations."""

    def test_engine_init_with_defaults(self):
        engine = QuoteExtractorEngine()
        assert engine.config is not None
        assert "part_number" in engine.config.anchors

    def test_engine_extract_file(self, single_modern_quote_file: Path):
        engine = QuoteExtractorEngine()
        doc = engine.extract_file(single_modern_quote_file)
        assert isinstance(doc, QuoteDocument)
        assert doc.quote_id == "Q1001"
        assert doc.total_items == 1
        assert doc.items[0].part_number == "2202046213000"

    def test_engine_extract_directory(self, batch_directory_fixture: Path):
        engine = QuoteExtractorEngine()
        docs = engine.extract_directory(str(batch_directory_fixture), recursive=True)
        assert len(docs) >= 4
        # Verify no lock files were processed
        assert not any(d.source_file.startswith("~$") for d in docs)
